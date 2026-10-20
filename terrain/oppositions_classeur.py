"""Import du suivi fourni : dossiers distincts et opposants multiples conservés."""
import hashlib
import json
import re
from datetime import date, datetime
from django.db import transaction
from .models import Opposition
from .imports import en_date


def lire_classeur(fichier):
    from openpyxl import load_workbook
    fichier.seek(0)
    wb=load_workbook(fichier,read_only=True,data_only=True)
    try:
        if 'SUIVI DES OPPOSITIONS' not in wb.sheetnames:return None
        sources={};details={};lignes=0
        for nom in ('SUIVI DES OPPOSITIONS','DETAILS OPPOSITIONS'):
            if nom not in wb.sheetnames:continue
            rows=iter(wb[nom].values);entetes=next(rows)
            for row in rows:
                numero=str(row[0] or '').strip()
                if not re.fullmatch(r'\d{3}-\d{3}-\d{5,12}',numero):continue
                valeurs={str(k).strip(): (v.isoformat() if isinstance(v,(date,datetime)) else (None if v in ('#N/A','#VALUE!') else v)) for k,v in zip(entetes[:19],row[:19]) if k}
                lignes+=1
                if nom=='SUIVI DES OPPOSITIONS':sources[numero]=valeurs
                else:
                    groupe=details.setdefault(numero,[])
                    if valeurs not in groupe:groupe.append(valeurs)
        dossiers=[]
        for numero in sorted(set(sources)|set(details)):
            enfants=details.get(numero,[])
            suivi=sources.get(numero)
            base=suivi or enfants[0]
            statut=str(base.get('STATUT OPPOSITION') or '').strip()
            if suivi is None and not all(str(d.get('STATUT OPPOSITION') or '').strip().casefold()=='opposition levée' for d in enfants):statut=''
            ref=hashlib.sha256(('suivi-oppositions:'+numero).encode()).hexdigest()
            dossiers.append(dict(source_reference=ref,num_demande=numero,
                departement=str(base.get('DEPARTEMENT') or '')[:40],
                sous_prefecture=str(base.get('SOUS-PREFECTURE') or '')[:100],
                village=str(base.get('VILLAGE') or '')[:100],
                code_parcelle=str(base.get('CODE DE PARCELLE') or '')[:60],
                demandeur=str(base.get('NOM REQUERANT') or '')[:160],
                opposant=str(base.get("NOM DE L'OPPOSANT") or '')[:160],
                motif=str(base.get("MOTIF DE L'OPPOSITION") or ''),type_litige='autre',
                date_opposition=en_date(base.get("DATE DE L'OPPOSITION")),
                statut='levee' if statut.casefold()=='opposition levée' else 'non_renseigne',
                source='import',source_metadata=dict(suivi=base,details=enfants,ce=base.get('CE') or '',publicite=base.get('STATUT PUBLICITE') or '',numero_atypique=not bool(re.fullmatch(r'\d{3}-\d{3}-\d{6}',numero)),observation=base.get('OBSERVATION') or '',surface=base.get('SURFACE (ha)'),surface_levee=base.get('SUPERFICIE LEVE'),fichier=getattr(fichier,'name',''))))
        if not dossiers:raise ValueError('Aucun dossier exploitable : les données existantes sont conservées.')
        return dossiers,lignes
    finally:
        wb.close();fichier.seek(0)


def importer_classeur(fichier,remplacer=False):
    if not str(getattr(fichier,'name','')).lower().endswith('.xlsx'):return None
    resultat=lire_classeur(fichier)
    if resultat is None:return None
    dossiers,lignes=resultat
    return enregistrer_dossiers(dossiers,lignes,remplacer)


def enregistrer_dossiers(dossiers,lignes=0,remplacer=False):
    with transaction.atomic():
        if remplacer:Opposition.objects.filter(source='import').delete()
        existants={o.source_reference:o for o in Opposition.objects.filter(source_reference__in=[d['source_reference'] for d in dossiers])}
        crees=[];maj=[]
        for d in dossiers:
            obj=existants.get(d['source_reference'])
            if obj is None:crees.append(Opposition(**d))
            else:
                for k,v in d.items():setattr(obj,k,v)
                maj.append(obj)
        Opposition.objects.bulk_create(crees,batch_size=200)
        if maj:Opposition.objects.bulk_update(maj,[k for k in dossiers[0] if k!='source_reference'],batch_size=200)
    return dict(cree=len(crees),actualise=len(maj),ignore=0,lignes=lignes,dossiers=len(dossiers),sans_date=sum(d['date_opposition'] is None for d in dossiers),sans_statut=sum(d['statut']=='non_renseigne' for d in dossiers))
